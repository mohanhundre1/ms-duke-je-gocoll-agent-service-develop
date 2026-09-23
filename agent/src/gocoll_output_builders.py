"""Build the ``Ctrl vs WF`` and ``JE Flags`` payloads for the shared renderer.

These are *data* builders, not writers: the shared ERP renderer
(``erp_type="gocoll"``) turns the two lists produced here into extra tabs
alongside ``Sheet1`` + ``PPSFGL05`` in the single JE workbook. Keeping the
rendering in the shared service means the GOCOLL agent owns no openpyxl code.

``Ctrl vs WF`` mirrors the A/B/C tie-out (coded vs WF vs Treasury); ``JE Flags``
consolidates the reconciliation gaps and every validation-check failure into one
analyst review list.
"""

from __future__ import annotations

import logging
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from src.models.enums import CheckName
from src.utils import strip_dot_zero

logger = logging.getLogger(__name__)

# Flag types that represent reconciliation results - surfaced in Ctrl_vs_WF tab,
# excluded from the GO Form / Validation Flags tabs.
RECON_FLAG_TYPES: frozenset[str] = frozenset(
    ("reconciliation_difference", "wf_missing", "return_item", "reconciliation_review", "missing_coding")
)

# SOP-prescribed analyst next actions per flag, surfaced in the review tab.
# Grounded in JA-SOP-01 GO Collections (Additional Comments, paras 202-219).
_MARBS_EMAIL = "MISC@duke-energy.com"

_NEXT_ACTIONS: dict[str, str] = {
    "missing_go_form_lines": (
        "Check has a bank/check amount but no GO-form coded distribution lines. "
        "Obtain the missing GO form or contact MARBS/user before posting."
    ),
    "reconciliation_difference": (
        "Investigate the coded-vs-WF gap; WF (bank) is authoritative. Correct the "
        "coded amount, or document a return item and add a Support-tab note."
    ),
    "wf_missing": (
        "Batch not found in the WF report. Obtain the correct WF report for this "
        "period (or confirm the batch with Treasury) before posting."
    ),
    "return_item": "No action required: gap explained by a BAI 566 return item.",
    "otc_deposit": (
        "OTC deposit found in the Treasury Report. The journal includes balanced "
        "OTC cash and negative detail lines. Before posting, confirm the accounting "
        "with Tim Coffey, complete the blank negative detail line, and "
        "retain the confirmation in support."
    ),
    "validation_tab": (
        "Validation issue: confirm the codeblock via the Derivation setup. If valid, "
        "add the only rejected value to its corresponding Validation Tab column/dropdown; "
        "if not valid, email the user (attach check/GO-form/error screenshots) and use "
        "Suspense (>$5,000) or Miscellaneous."
    ),
    "balance": (
        "Journal does not balance: correct line amounts so the Monetary Amount Total "
        "= 0 before upload."
    ),
    "control_totals": (
        "Batch does not net to zero: correct the batch's cash/detail lines."
    ),
    "line_count": (
        "Line count differs from expected: verify no lines were dropped or added."
    ),
}

NO_GO_FORM_ACTION = (
    f"No GO form: email MARBS ({_MARBS_EMAIL}) to confirm the codeblock or exclude "
    f"the payment. If MARBS is unsure, use Suspense (>$5,000) or Miscellaneous."
)
ILLEGIBLE_FORM_ACTION = (
    "Illegible/partial GO form: contact the user who submitted it. If unresolved, "
    "use Suspense (>$5,000) or Miscellaneous (Support tab)."
)

BATCH_SEQ_RE = re.compile(r"\bBatch\s+(?P<batch>\S+)(?:\s+seq\s+(?P<seq>\d+))?", re.IGNORECASE)
MISSING_DIMS_RE = re.compile(
    r"\bBatch\s+(?P<batch>\S+)\s+seq\s+(?P<seq>\d+):\s+missing\s+(?P<missing>.+?)\s+\(left blank\)",
    re.IGNORECASE,
)
NO_GO_FORM_RE = re.compile(
    r"\bBatch\s+(?P<batch>\S+)\s+seq\s+(?P<seq>\d+):\sno GO form matched",
    re.IGNORECASE,
)


def _next_action(flag_type: str, detail: str) -> str:
    """Map a flag to the SOP-prescribed analyst next action."""
    if flag_type in ("codeblock_completeness", "missing_go_block", "missing_code_block_fields"):
        return (
            NO_GO_FORM_ACTION
            if "no go form" in (detail or "").lower()
            else ILLEGIBLE_FORM_ACTION
        )
    return _NEXT_ACTIONS.get(flag_type, "Review and resolve per SOP.")


def _money(value: Decimal | float | int | None) -> float | str:
    """Decimal/number -> float for the money-formatted cells; ``None`` -> blank."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    return ""


def _check_display(value: Any) -> str:
    """Keep only real numeric check identifiers in analyst-facing flags."""
    text = str(value or "").strip()
    return text if text.isdigit() else ""


def _batch_seq_from_detail(detail: str) -> tuple[str, str]:
    """Extract batch and line sequence from standard validation failure text."""
    match = _BATCH_SEQ_RE.search(detail or "")
    if not match:
        return "", ""
    return match.group("batch") or "", match.group("seq") or ""


def _assembly_check_number(pipeline: Any, batch: str, seq: str) -> str:
    """Resolve a validation failure's displayed batch/sequence to its check."""
    assembly = getattr(pipeline, "assembly", None)
    journal = getattr(assembly, "journal_entry", None)
    for line in getattr(journal, "all_lines", []) or []:
        if (
            str(getattr(line, "batch_number", "") or "") == str(batch)
            and str(getattr(line, "line_seq", "") or "") == str(seq)
        ):
            return _check_display(getattr(line, "check_number", ""))
    return ""


def _sequence_summary(seqs: list[str]) -> str:
    """Collapse sequence values into a compact list/range string."""
    nums = sorted(int(s) for s in seqs if str(s).isdigit())
    if not nums:
        return ""
    ranges: list[str] = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        ranges.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = n
    ranges.append(str(start) if start == prev else f"{start}-{prev}")
    return ", ".join(ranges)


def _append_codeblock_flags(flags: list[dict[str, Any]], failures: list[str], severity: str) -> None:
    """Group repetitive codeblock completeness failures by batch/reason."""
    grouped: dict[tuple[str, str, str], list[str]] = {}
    passthrough: list[str] = []

    for failure in failures:
        no_form = _NO_GO_FORM_RE.search(failure)
        if no_form:
            grouped.setdefault(
                (no_form.group("batch"), "no_go_form", "No GO form matched - posted with a blank code block"),
                [],
            ).append(no_form.group("seq"))
            continue

        missing = _MISSING_DIMS_RE.search(failure)
        if missing:
            reason = f"missing {missing.group('missing')} (left blank)"
            grouped.setdefault((missing.group("batch"), "missing_dimensions", reason), []).append(
                missing.group("seq")
            )
            continue

        passthrough.append(failure)

    for (batch, subtype, reason), seqs in sorted(grouped.items()):
        seq_summary = _sequence_summary(seqs)
        detail = f"Batch {batch} seq {seq_summary}: {reason}"
        if len(seqs) > 1:
            detail += f" ({len(seqs)} lines)"
        flag_type = "missing_go_block" if subtype == "no_go_form" else "missing_code_block_fields"
        flags.append(
            {
                "batch": batch,
                "seq": seq_summary,
                "check": "",
                "type": flag_type,
                "detail": detail,
                "amount": "",
                "confidence": "",
                "severity": severity,
                "next_action": _next_action(flag_type, detail),
                "subtype": subtype,
            }
        )

    for failure in passthrough:
        batch, seq = _batch_seq_from_detail(failure)
        flags.append(
            {
                "batch": batch,
                "seq": seq,
                "check": "",
                "type": "codeblock_completeness",
                "detail": failure,
                "amount": "",
                "confidence": "",
                "severity": severity,
                "next_action": _next_action("codeblock_completeness", failure),
            }
        )


_CB_FIELDS = ("business_unit", "account", "resource_type", "operating_unit",
              "resp_center", "project", "activity_id", "process",
              "location", "product", "affiliate", "alloc_pool")


def _cb_field(cb: Any, field: str) -> str:
    return getattr(cb, field, "") if cb is not None else ""


def build_extracted_data_rows(pipeline: Any) -> list[dict[str, Any]]:
    """One row per extracted transaction: batch, check, seq, GL dims, amount."""
    extraction = getattr(pipeline, "extraction", None)
    if extraction is None:
        return []
    rows: list[dict[str, Any]] = []
    for batch in getattr(extraction, "batches", []) or []:
        batch_no = getattr(batch, "batch_number", "")
        for check in getattr(batch, "checks", []) or []:
            check_no = str(getattr(check, "check_number", "") or "")
            seq_no = str(getattr(check, "sequence_number", "") or "")
            for ln in getattr(check, "lines", []) or []:
                cb = getattr(ln, "code_block", None)
                rows.append(
                    {
                        "batch_number": batch_no,
                        "check_number": check_no or str(getattr(ln, "check_number", "") or ""),
                        "sequence_number": seq_no,
                        "amount": _money(getattr(ln, "amount", None)),
                        **{f: _cb_field(cb, f) for f in _CB_FIELDS},
                        "confidence": round(float(getattr(ln, "confidence", 1.0) or 1.0), 2),
                    }
                )
    return rows


def _read_source_sheet_rows(source: str, sheet_hint: str) -> list[list[Any]]:
    """Read a source worksheet without replacing formulas with cached values."""
    if not source:
        return []
    try:
        from pathlib import Path as _Path

        import openpyxl as _xl

        path = _Path(source)
        if not path.exists():
            return []
        wb = _xl.load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = next(
                (name for name in wb.sheetnames if sheet_hint.lower() in name.lower()),
                wb.sheetnames[0],
            )
            ws = wb[sheet]
            # Excel can persist a <dimension> spanning the whole grid even when
            # only a few rows hold data; read-only mode trusts it and would
            # otherwise materialise millions of empty cells.
            ws.reset_dimensions = True
            rows: list[list[Any]] = [
                [cell.value for cell in row]
                for row in ws.iter_rows()
            ]
            finally:
                wb.close()
            while rows and all(cell in (None, "") for cell in rows[-1]):
                rows.pop()
            return rows
        except Exception:  # noqa: BLE001 - support-tab copy is best-effort
            logger.warning("Failed to copy source sheet from %s", source, exc_info=True)
            return []


def build_treasury_rows(pipeline: Any) -> list[list[Any]]:
    """Copy the input Treasury sheet while preserving its formulas."""
    treasury = getattr(pipeline, "treasury_report", None)
    source = getattr(treasury, "source_file", "") if treasury is not None else ""
    return _read_source_sheet_rows(source, "treasury")


def build_wf_transactions_rows(pipeline: Any) -> list[list[Any]]:
    """Copy the input WF Transactions sheet while preserving its formulas."""
    report = getattr(pipeline, "wf_report", None)
    source = getattr(report, "source_file", "") if report is not None else ""
    return _read_source_sheet_rows(source, "transaction")


def build_support_sources(
    validation_source: str,
    wf_source: str,
    treasury_source: str,
    *,
    fm_url_map: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    """Ship the classified source workbooks for the renderer to merge itself.

    Each entry carries only its canonical input FilePart URL. Source workbook
    bytes must never be copied into the AZA result because the next workflow
    step traverses Dapr's bounded router request body.

    Deliberately not named ``source_file_contents``: that key is this agent's own
    inbound file contract and the ERP agent copies it straight into pipeline_data.
    """
    sources = [
        ("Validation Tab", validation_source, "validation"),
        ("WF Transactions", wf_source, "transaction"),
        ("Treasury", treasury_source, "treasury"),
    ]
    contents: list[dict[str, str]] = []
    fm_url_map = fm_url_map or {}
    for sheet_title, source, hint in sources:
        fm_url = fm_url_map.get(hint, "")
        if not source or not Path(source).exists():
            if fm_url:
                # Local file absent but FM URL available - renderer downloads.
                contents.append({"sheet_title": sheet_title, "sheet_hint": hint,
                                 "filename": "", "fm_url": fm_url})
                continue
            path = Path(source)
            entry: dict[str, str] = {"sheet_title": sheet_title, "sheet_hint": hint,
                                     "filename": path.name}
            if not fm_url:
                logger.error("Missing File Manager URL for GOCOLL %s source %s", hint, path.name)
                continue
            entry["fm_url"] = fm_url
            contents.append(entry)

    logger.info(
        "GOCOLL support sources: %d workbook FM URL(s)",
        len(contents),
    )
    return contents


def build_validation_lists(pipeline: Any) -> dict[str, list[str]]:
    """Valid dimension values per Sheet1 column, for the renderer's dropdowns.

    Keyed by eFIS Sheet1 column key so the renderer can attach an Excel list
    data-validation. Empty when no Validation Tab master was loaded.
    """
    master = getattr(pipeline, "validation_master", None)
    if master is None or not getattr(master, "loaded", False):
        return {}

    def _sorted(values: Any) -> list[str]:
        return sorted((str(v).strip() for v in (values or []) if str(v).strip()))

    lists = {
        "J_line_bus_unit": _sorted(getattr(master, "business_units", ())),
        "K_account": _sorted(getattr(master, "active_accounts", ())),
        "L_resource_type": _sorted(getattr(master, "resource_types", ())),
        "M_oper_unit": _sorted(getattr(master, "operating_units", ())),
        "N_resp_center": _sorted(getattr(master, "resp_centers", ())),
        "O_project": _sorted(getattr(master, "projects", ())),
        "P_activity_id": _sorted(getattr(master, "activities", ())),
        "Q_process": _sorted(getattr(master, "processes", ())),
        "R_location": _sorted(getattr(master, "locations", ())),
    }
    return {key: values for key, values in lists.items() if values}


_SHEET1_DROPDOWN_KEYS = (
    "J_line_bus_unit",
    "K_account",
    "L_resource_type",
    "M_oper_unit",
    "N_resp_center",
    "O_project",
    "P_activity_id",
    "Q_process",
    "R_location",
)


def _normalize_dropdown_value(key: str, value: Any) -> str:
    text = strip_dot_zero(str(value).strip()) if value is not None else ""
    if key == "K_account" and text.isdigit():
        return text.zfill(7)
    return text


def clear_values_outside_dropdowns(
    rows: list[dict[str, Any]],
    validation_lists: dict[str, list[str]],
) -> list[dict[str, Any]]:
    """Blank dimensions that cannot be selected from their Sheet1 dropdown."""
    if not validation_lists:
        return rows

    allowed_by_key = {
        key: {_normalize_dropdown_value(key, value) for value in validation_lists.get(key, [])}
        for key in _SHEET1_DROPDOWN_KEYS
        if validation_lists.get(key)
    }
    if not allowed_by_key:
        return rows

    cleaned_rows: list[dict[str, Any]] = []
    for row in rows:
        cleaned = dict(row)
        for key, allowed in allowed_by_key.items():
            raw = cleaned.get(key, "")
            if raw not in (None, "") and _normalize_dropdown_value(key, raw) not in allowed:
                cleaned[key] = ""
        cleaned_rows.append(cleaned)
    return cleaned_rows


def build_ctrl_vs_wf_rows(pipeline: Any) -> list[dict[str, Any]]:
    """One row per reconciled batch: coded (A) vs WF (B) vs Treasury (C)."""
    recon = getattr(pipeline, "reconciliation", None)
    if recon is None or not recon.batches:
        return []
    rows: list[dict[str, Any]] = []
    for b in sorted(recon.batches, key=lambda x: x.batch_number):
        rows.append(
            {
                "batch": b.batch_number,
                "coded": _money(b.coded_amount),
                "wf": _money(b.wf_amount),
                "treasury": _money(b.treasury_amount),
                "return_items": _money(b.return_items),
                "diff_coded_wf": _money(b.diff_coded_wf),
                "diff_wf_treasury": _money(b.diff_wf_treasury),
                "status": b.status,
                "detail": b.detail or "",
            }
        )
    return rows


# Per-check coded-vs-WF tie-out tolerance (rounding only).
_CHECK_TIE_TOL = Decimal("0.01")


def build_ctrl_vs_wf_check_rows(pipeline: Any) -> list[dict[str, Any]]:
    """Per-check tie-out: each extracted check's coded line total vs the WF
    report's per-check amounts, matched within a batch.

    Finer than the per-batch A/B/C tie-out: a single check misread that a
    batch-total comparison would mask (two checks off by offsetting amounts, or a
    dropped/duplicated check) surfaces here. Matching is by AMOUNT within the
    batch, so the check number is never required - the SME data-privacy redaction
    of check numbers does not affect this. WF is authoritative per the SOP
    ("the comparison is with the bank itself and that information is certain").

    Returns [] when no WF report was supplied (no bank record to tie to).
    """
    wf = getattr(pipeline, "wf_report", None)
    extraction = getattr(pipeline, "extraction", None)
    if wf is None or not getattr(wf, "loaded", False) or extraction is None:
        return []

    # Prefer the split-aware, correction-applied rows computed before assembly.
    stored = getattr(extraction, "check_recon_rows", None)
    if stored:
        return stored

    rows: list[dict[str, Any]] = []
    for batch in getattr(extraction, "batches", []) or []:
        batch_no = getattr(batch, "batch_number", "")
        # WF per-check amounts for this batch (the bank's record).
        wf_pool = [
            abs(t.check_amount) for t in wf.transactions_for(batch_no) if t.check_amount
        ]
        checks = list(getattr(batch, "checks", []) or [])

        # Uncoded checks (deposited amount but no GO-form lines) are reported in
        # the GO Form Flags tab; consume their WF amount here so they are not
        # double-counted as an unmatched coded check below.
        for check in checks:
            if getattr(check, "lines", None):
                continue
            amt = abs(getattr(check, "check_amount", Decimal("0")) or Decimal("0"))
            j = next(
                (i for i, a in enumerate(wf_pool) if abs(a - amt) <= _CHECK_TIE_TOL),
                None,
            )
            if j is not None:
                wf_pool.pop(j)

        # Coded checks: tie the coded line total to a WF check amount by value.
        for check in checks:
            lines = getattr(check, "lines", None) or []
            if not lines:
                continue
            coded_sum = sum((abs(ln.amount) for ln in lines), Decimal("0"))
            check_no = _check_display(getattr(check, "check_number", ""))
            j = next(
                (i for i, a in enumerate(wf_pool) if abs(a - coded_sum) <= _CHECK_TIE_TOL),
                None,
            )
            if j is not None:
                wf_amt = wf_pool.pop(j)
                rows.append(
                    {
                        "batch": batch_no,
                        "check": check_no,
                        "coded": _money(coded_sum),
                        "wf": _money(wf_amt),
                        "diff": _money(coded_sum - wf_amt),
                        "status": "tie",
                        "detail": "Coded check total ties to a WF check amount",
                    }
                )
            else:
                rows.append(
                    {
                        "batch": batch_no,
                        "check": check_no,
                        "coded": _money(coded_sum),
                        "wf": "",
                        "diff": "",
                        "status": "no_wf_match",
                        "detail": (
                            f"Coded total {coded_sum} matches no WF check amount in "
                            f"batch {batch_no} - verify the coded amount against the bank"
                        ),
                    }
                )

        # WF checks with no coded match left over - the bank shows a check we did
        # not code (a check may be missing or its amount misread).
        for amt in wf_pool:
            rows.append(
                {
                    "batch": batch_no,
                    "check": "",
                    "coded": "",
                    "wf": _money(amt),
                    "diff": "",
                    "status": "missing_coding",
                    "detail": (
                        f"WF check amount {amt} in batch {batch_no} has no matching "
                        f"coded check - a check may be missing or misread"
                    ),
                }
            )

    return rows


def build_flags_rows(pipeline: Any) -> list[dict[str, Any]]:
    """Consolidated review flags: reconciliation gaps + validation failures."""
    flags: list[dict[str, Any]] = []

    treasury_report = getattr(pipeline, "treasury_report", None)
    otc_amount = getattr(treasury_report, "total_over_the_counter", Decimal("0"))
    if otc_amount and otc_amount != Decimal("0"):
        detail = (
            f"Treasury Report contains an over-the-counter (OTC) deposit totaling "
            f"{otc_amount}. Balanced OTC lines were added to Sheet1, but the "
            f"negative detail line has blank accounting dimensions pending a "
            f"confirmed codeblock."
        )
        flags.append(
            {
                "batch": "",
                "seq": "",
                "check": "",
                "type": "otc_deposit",
                "detail": detail,
                "amount": _money(abs(otc_amount)),
                "confidence": "",
                "severity": "review",
                "next_action": _next_action("otc_deposit", detail),
            }
        )

    extraction = getattr(pipeline, "extraction", None)

    # WF amounts already claimed by shared-form split/corrected blocks - used
    # to suppress missing_go_form_lines flags for uncoded check-front pages
    # whose amounts are coded on a multi-check GO form elsewhere in the batch.
    recon_rows = getattr(extraction, "check_recon_rows", []) or []
    _covered_wf: set[float] = {
        float(r["wf"])
        for r in recon_rows
        if r.get("status") == "tie" and r.get("wf")
    }

    missing_check_groups: dict[str, dict[str, Any]] = {}
    for batch in getattr(extraction, "batches", []) or []:
        for check in getattr(batch, "checks", []) or []:
            amount = getattr(check, "check_amount", None)
            has_amount = bool(amount and amount != Decimal("0"))
            try:
                _amt_f = float(abs(amount)) if amount else 0.0
            except (TypeError, ValueError):
                _amt_f = 0.0
            _covered = any(abs(_amt_f - c) < 0.02 for c in _covered_wf)
            if has_amount and not (getattr(check, "lines", None) or []) and not _covered:
                batch_no = getattr(batch, "batch_number", "")
                group = missing_check_groups.setdefault(
                    batch_no,
                    {"seqs": [], "checks": [], "amount": Decimal("0")},
                )
                seq = getattr(check, "sequence_number", "") or ""
                if seq:
                    group["seqs"].append(seq)
                group["checks"].append(_check_display(getattr(check, "check_number", "")) or "(unknown)")
                if amount is not None:
                    group["amount"] += abs(amount)

    for batch_no, group in sorted(missing_check_groups.items()):
        checks = group["checks"]
        known_checks = [c for c in checks if c != "(unknown)"]
        unknown_count = len(checks) - len(known_checks)
        check_bits = known_checks[:8]
        if unknown_count:
            check_bits.append(f"(unknown) x{unknown_count}")
        if len(known_checks) > 8:
            check_bits.append(f"+{len(known_checks) - 8} more")
        seq_summary = _sequence_summary(group["seqs"])
        detail = (
            f"Batch {batch_no}: {len(checks)} check(s) have check amount but no "
            f"GO-form coded distribution lines"
        )
        if check_bits:
            detail += f"; checks: {', '.join(check_bits)}"
        flags.append(
            {
                "batch": batch_no,
                "seq": seq_summary,
                "check": ", ".join(check_bits),
                "type": "missing_go_form_lines",
                "detail": detail,
                "amount": _money(group["amount"]),
                "confidence": "",
                "severity": "review",
                "next_action": _next_action("missing_go_form_lines", detail),
            }
        )

    recon = getattr(pipeline, "reconciliation", None)
    if recon is not None:
        for b in recon.unexplained:
            diff = b.diff_coded_wf
            detail = b.detail or f"Coded {b.coded_amount} vs WF {b.wf_amount} - unexplained gap"
            flags.append(
                {
                    "batch": b.batch_number,
                    "seq": "",
                    "check": "",
                    "type": "reconciliation_difference",
                    "detail": detail,
                    "amount": _money(abs(diff)) if diff is not None else "",
                    "confidence": "",
                    "severity": "review",
                    "next_action": _next_action("reconciliation_difference", detail),
                }
            )
        for b in recon.wf_missing:
            detail = b.detail or f"Batch {b.batch_number} not found in the WF report - no bank record"
            flags.append(
                {
                    "batch": b.batch_number,
                    "seq": "",
                    "check": "",
                    "type": "wf_missing",
                    "detail": detail,
                    "amount": _money(b.coded_amount),
                    "confidence": "",
                    "severity": "review",
                    "next_action": _next_action("wf_missing", detail),
                }
            )
        for b in recon.return_item_batches:
            detail = b.detail or f"A-vs-B gap explained by {b.return_items} return item(s)"
            flags.append(
                {
                    "batch": b.batch_number,
                    "seq": "",
                    "check": "",
                    "type": "return_item",
                    "detail": detail,
                    "amount": _money(b.return_items),
                    "confidence": "",
                    "severity": "info",
                    "next_action": _next_action("return_item", detail),
                }
            )

    # Every other validation-check failure (the "reconciliation" check is already
    # represented above as structured rows, so skip it to avoid duplicates).
    processing_result = getattr(pipeline, "processing_result", None)
    for chk in getattr(processing_result, "checks", []) or []:
        if chk.check_name in (CheckName.RECONCILIATION, CheckName.DERIVA_GUARD):
            continue
        status = str(getattr(chk, "status", "value", chk.status)).upper()
        severity = "fail" if status == "FAIL" else "review"
        if chk.check_name == CheckName.CODEBLOCK_COMPLETENESS:
            _append_codeblock_flags(flags, list(chk.failures or []), severity)
            continue
        for failure in chk.failures or []:
            batch, seq = _batch_seq_from_detail(failure)
            flags.append(
                {
                    "batch": batch,
                    "seq": seq,
                    "check": _assembly_check_number(pipeline, batch, seq),
                    "type": chk.check_name,
                    "detail": failure,
                    "amount": "",
                    "confidence": "",
                    "severity": severity,
                    "next_action": _next_action(chk.check_name, failure),
                }
            )

    return flags