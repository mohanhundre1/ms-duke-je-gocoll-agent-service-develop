"""GOColl input detection & mandatory-file validation.

The GOColl JE run requires these input feeds:

* **Validation Tab** workbook - PeopleSoft dimension master (analyst-provided)
* **WF Transaction Report** - Wells Fargo lockbox transaction report
* **Batch PDFs** - one or more lockbox batch PDFs (>= 1)

The **Treasury Report** (bank "Total Bank Transactions" feed) is *optional*: when
provided it supplies return-item (BAI 566) and lockbox (BAI 115) controls; when
absent, reconciliation still runs against WF and simply cannot explain a gap with
a Treasury return item.

Each incoming file is classified by **filename keyword first, then file content**
(sheet names / markers). A single "combined" GOColllections workbook that carries
the Validation Tab, Treasury and WF sheets together satisfies all three workbook
feeds at once.

Before extraction the caller must assert every mandatory feed is present and
non-empty; if any is missing or empty the run is a **hard stop** (no partial JE).
A missing Treasury feed is not a hard stop.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path

logger = logging.getLogger(__name__)

_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".xltx", ".xltm"}

class GoCollFileType(StrEnum):
    VALIDATION_TAB = "validation_tab"
    TREASURY = "treasury"
    WF_REPORT = "wf_report"
    BATCH_PDF = "batch_pdf"
    UNKNOWN = "unknown"

# Filename keyword aliases - checked case-insensitively; first hit wins.
# Kept specific enough so a bare "GoColllections_*.xlsx" combined workbook
# falls through to content detection instead of being mislabelled.
_FILENAME_RULES: list[tuple[GoCollFileType, tuple[str, ...]]] = [
    (GoCollFileType.VALIDATION_TAB, (
        "validation tab", "validation_tab", "valid tab", "valid_tab",
        "validationtab", "gocoll validation", "gocoll_validation",
    )),
    (GoCollFileType.WF_REPORT, (
        "wf transaction", "wf_transaction", "wf transactions", "wf_transactions",
        "wells fargo", "wellsfargo", "wf report", "wf_report",
        "lockbox report", "lockbox_report",
    )),
    (GoCollFileType.TREASURY, (
        "treasury", "total bank", "bank transactions", "bank_transactions",
        "without groups", "duke total bank",
    )),
]

def _detect_by_filename(path: Path) -> GoCollFileType | None:
    name = path.name.lower()
    for ftype, keywords in _FILENAME_RULES:
        if any(k in name for k in keywords):
            return ftype
    # Any PDF is a batch PDF - no other PDF input type exists for GOColl
    if path.suffix.lower() == ".pdf":
        return GoCollFileType.BATCH_PDF
    return None

def _sheet_cells(ws, max_row: int = 3) -> list[str]:
    """Flatten the first *max_row* rows of a worksheet into lowercased cell strings."""
    values: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, values_only=True):
        for cell in row:
            if cell is not None:
                values.append(str(cell).lower())
    return values

def _matches(cells: list[str], *aliases: str) -> bool:
    """True if any cell contains any alias as a substring (case already lowered)."""
    return any(alias in cell for cell in cells for alias in aliases)

# -- Content marker aliases -------------------------------------------------
# Each tuple lists all known name variants for that column header/title.

_VT_BU = ("business unit", "bus unit", "bu ", "businessunit")
_VT_ACCOUNT = ("account number", "account no", "acct number", "acct no", "account nbr")
_VT_RESTYPE = ("resource type", "res type", "resourcetype", "resource typ")

_TR_TITLE = ("duke total bank transactions", "total bank transactions without groups",
             "total bank transactions")
_TR_BAI = ("bai code", "bai_code", "bai ")
_TR_DEBIT = ("debit",)
_TR_CREDIT = ("credit",)

_WF_BATCH = ("batch number", "batch no", "batch nbr", "batch_number")
_WF_CHECK = ("check amount", "check amt", "chk amount", "chk amt", "check_amount")
_WF_LOCKBOX = ("lockbox", "lock box", "lock_box", "lbx")

def _is_validation_tab_sheet(cells: list[str]) -> bool:
    return (
        _matches(cells, *_VT_BU)
        and _matches(cells, *_VT_ACCOUNT)
        and _matches(cells, *_VT_RESTYPE)
    )

def _is_treasury_sheet(cells: list[str]) -> bool:
    return _matches(cells, *_TR_TITLE) or (
        _matches(cells, *_TR_BAI)
        and (_matches(cells, *_TR_DEBIT) or _matches(cells, *_TR_CREDIT))
    )

def _is_wf_report_sheet(cells: list[str]) -> bool:
    return (
        _matches(cells, *_WF_BATCH)
        and _matches(cells, *_WF_CHECK)
        and _matches(cells, *_WF_LOCKBOX)
    )

def _detect_by_content(path: Path) -> set[GoCollFileType]:
    """Infer type(s) from file content. A combined workbook may return several."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return {GoCollFileType.BATCH_PDF}
    if suffix not in _EXCEL_SUFFIXES:
        return set()

    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - unreadable workbook -> unknown
        logger.warning("gocoll_inputs: could not open %s for content detection: %s", path.name, exc)
        return set()

    types: set[GoCollFileType] = set()
    try:
        for sheet in wb.sheetnames:
            s = sheet.lower()
            # Exact known sheet names from combined workbooks - fast path
            if s == "validation tab":
                types.add(GoCollFileType.VALIDATION_TAB)
                continue
            if s == "treasury report":
                types.add(GoCollFileType.TREASURY)
                continue
            if s == "wf transactions report":
                types.add(GoCollFileType.WF_REPORT)
                continue
            # Skip sheets that belong to GOColl output files or helper tabs
            if s in ("flagged for review", "ctrl.vs wf report", "ctrl vs wf",
                     "deriva listing", "additional support", "otc control",
                     "mtd check", "instructions", "review checklist", "tables"):
                continue
            # Cell-content scan for generic sheet names (e.g. "Sheet1")
            cells = _sheet_cells(wb[sheet])
            if _is_validation_tab_sheet(cells):
                types.add(GoCollFileType.VALIDATION_TAB)
            if _is_treasury_sheet(cells):
                types.add(GoCollFileType.TREASURY)
            if _is_wf_report_sheet(cells):
                types.add(GoCollFileType.WF_REPORT)
    finally:
        wb.close()
    return types

@dataclass
class InputSet:
    """Resolved GOColl input feeds."""

    validation_tab: Path | None = None
    treasury: Path | None = None
    wf_report: Path | None = None
    batch_pdfs: list[Path] = field(default_factory=list)
    unknown: list[Path] = field(default_factory=list)

    def missing_mandatory(self) -> list[str]:
        # Treasury is optional (return-item/lockbox controls); its absence only
        # degrades reconciliation, it does not block the run.
        missing: list[str] = []
        if self.validation_tab is None:
            missing.append(GoCollFileType.VALIDATION_TAB.value)
        if self.wf_report is None:
            missing.append(GoCollFileType.WF_REPORT.value)
        if not self.batch_pdfs:
            missing.append(GoCollFileType.BATCH_PDF.value)
        return missing

class GoCollInputError(ValueError):
    """Raised when mandatory inputs are missing or empty (hard stop)."""

def classify_inputs(paths: list[str | Path]) -> InputSet:
    """Classify each input file into the GOColl feed it satisfies.

    Filename keywords take precedence; on a miss the file's content (sheet names
    for workbooks, extension for PDFs) is used. A combined workbook that holds
    several feed sheets is assigned to each empty matching slot.
    """
    result = InputSet()

    def _assign(ftype: GoCollFileType, p: Path) -> None:
        if ftype is GoCollFileType.BATCH_PDF:
            if p not in result.batch_pdfs:
                result.batch_pdfs.append(p)
        elif ftype is GoCollFileType.VALIDATION_TAB and result.validation_tab is None:
            result.validation_tab = p
        elif ftype is GoCollFileType.TREASURY and result.treasury is None:
            result.treasury = p
        elif ftype is GoCollFileType.WF_REPORT and result.wf_report is None:
            result.wf_report = p

    for raw in paths:
        p = Path(raw)
        by_name = _detect_by_filename(p)
        if by_name is not None:
            _assign(by_name, p)
            continue
        content_types = _detect_by_content(p)
        if content_types:
            for ftype in content_types:
                _assign(ftype, p)
        else:
            result.unknown.append(p)

    return result

def _is_empty_file(path: Path) -> str | None:
    """Return an error string if the file is missing/empty, else None."""
    if not path.exists():
        return f"{path.name}: file not found"
    try:
        if path.stat().st_size == 0:
            return f"{path.name}: file is empty (0 bytes)"
    except OSError as exc:
        return f"{path.name}: cannot stat file ({exc})"

    if path.suffix.lower() in _EXCEL_SUFFIXES:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                has_rows = any(
                    any(cell is not None for cell in row)
                    for ws in wb.worksheets
                    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True)
                )
            finally:
                wb.close()
            if not has_rows:
                return f"{path.name}: workbook has no data rows"
        except Exception as exc:  # noqa: BLE001
            return f"{path.name}: workbook could not be read ({exc})"
    return None

def validate_inputs(input_set: InputSet) -> list[str]:
    """Return a list of hard-stop errors (missing mandatory feeds + empty files).

    An empty list means the inputs are complete and ready for extraction.
    """
    errors: list[str] = []

    missing = input_set.missing_mandatory()
    if missing:
        errors.append(
            "Missing mandatory input(s): " + ", ".join(missing)
        )

    to_check: list[Path] = []
    for feed in (input_set.validation_tab, input_set.treasury, input_set.wf_report):
        if feed is not None:
            to_check.append(feed)
    to_check.extend(input_set.batch_pdfs)

    seen: set[Path] = set()
    for path in to_check:
        if path in seen:
            continue
        seen.add(path)
        err = _is_empty_file(path)
        if err:
            errors.append(err)

    return errors

def classify_and_validate_inputs(paths: list[str | Path]) -> InputSet:
    """Classify inputs and hard-stop (raise) if any mandatory feed is missing/empty."""
    input_set = classify_inputs(paths)
    errors = validate_inputs(input_set)
    if errors:
        raise GoCollInputError(
            "GOColl input validation failed (hard stop): " + "; ".join(errors)
        )
    logger.info(
        "GOColl inputs OK: validation_tab=%s treasury=%s wf_report=%s batch_pdfs=%d unknown=%d",
        getattr(input_set.validation_tab, "name", None),
        getattr(input_set.treasury, "name", None),
        getattr(input_set.wf_report, "name", None),
        len(input_set.batch_pdfs),
        len(input_set.unknown),
    )
    return input_set